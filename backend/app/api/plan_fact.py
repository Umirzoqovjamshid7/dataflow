from collections import defaultdict
from calendar import monthrange
from datetime import date, datetime, time, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy.orm import Session

from app.api.deps import current_user, require_roles, tenant_query
from app.core.db import get_db
from app.models.models import Campaign, Lead, V5BusinessPlan

router = APIRouter(prefix="/api/plan-fact", tags=["plan-fact"])

class PlanUpdate(BaseModel):
    model_config = ConfigDict(extra="forbid")
    month: str = Field(pattern=r"^\d{4}-(0[1-9]|1[0-2])$")
    sales: Decimal = Field(ge=0, max_digits=18, decimal_places=2)
    revenue: Decimal = Field(ge=0, max_digits=18, decimal_places=2)

@router.put("")
def save_plan(data: PlanUpdate, user=Depends(require_roles("tenant_admin")), db: Session = Depends(get_db)):
    if not user.tenant_id:
        raise HTTPException(409, "Kompaniya akkaunti talab qilinadi")
    for metric, target in (("sales", data.sales), ("revenue", data.revenue)):
        row = db.query(V5BusinessPlan).filter_by(tenant_id=user.tenant_id, month=data.month, metric=metric, project="__total__").first()
        if row: row.target = target
        else: db.add(V5BusinessPlan(tenant_id=user.tenant_id, month=data.month, metric=metric, project="__total__", target=target))
    db.commit()
    return {"saved": True, "month": data.month, "sales": float(data.sales), "revenue": float(data.revenue)}


@router.get("/template")
def plan_template(_user=Depends(current_user)):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Plan-Fact"
    sheet.append(["month", "project", "sales_plan", "revenue_plan"])
    sheet.append([datetime.now(timezone.utc).strftime("%Y-%m"), "Asosiy loyiha", 25, 75000000])
    sheet.freeze_panes = "A2"
    for column, width in {"A": 14, "B": 32, "C": 18, "D": 20}.items():
        sheet.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="plan-fact-template.xlsx"'})


def normalized_header(value):
    return str(value or "").strip().lower().replace("-", "_").replace(" ", "_")


def parsed_month(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m")
    result = str(value or "").strip()[:7]
    try:
        datetime.strptime(result, "%Y-%m")
    except ValueError as exc:
        raise ValueError("oy YYYY-MM formatida bo‘lishi kerak") from exc
    return result


def nonnegative(value, field):
    try:
        result = Decimal(str(value if value not in (None, "") else 0))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{field} raqam bo‘lishi kerak") from exc
    if not result.is_finite():
        raise ValueError(f"{field} chekli raqam bo‘lishi kerak")
    if result < 0:
        raise ValueError(f"{field} manfiy bo‘lishi mumkin emas")
    if result > Decimal("9999999999999999.99"):
        raise ValueError(f"{field} juda katta")
    if result.quantize(Decimal("0.01")) != result:
        raise ValueError(f"{field} ko‘pi bilan 2 kasr xonasiga ega bo‘lishi kerak")
    return result


@router.post("/import")
async def import_plan(file: UploadFile = File(...), user=Depends(require_roles("tenant_admin")), db: Session = Depends(get_db)):
    if not user.tenant_id:
        raise HTTPException(409, "Kompaniya akkaunti talab qilinadi")
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(422, "Faqat .xlsx Excel fayli qabul qilinadi")
    content = await file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(413, "Excel fayli 5 MB dan oshmasligi kerak")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [normalized_header(value) for value in next(rows)]
    except Exception as exc:
        raise HTTPException(422, "Excel faylini o‘qib bo‘lmadi") from exc
    aliases = {
        "month": {"month", "oy", "месяц"},
        "project": {"project", "loyiha", "проект"},
        "sales": {"sales_plan", "savdo_rejasi", "savdo_reja", "plan", "reja"},
        "revenue": {"revenue_plan", "summa_rejasi", "summa_reja", "amount_plan"},
    }
    indexes = {key: next((headers.index(alias) for alias in names if alias in headers), None) for key, names in aliases.items()}
    if any(indexes[key] is None for key in aliases):
        raise HTTPException(422, "Ustunlar kerak: month, project, sales_plan, revenue_plan")
    imported, number = 0, 1
    try:
        for number, values in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in values):
                continue
            if imported >= 5000:
                raise ValueError("Bitta faylda 5000 qatordan oshmasligi kerak")
            month = parsed_month(values[indexes["month"]])
            project = str(values[indexes["project"]] or "").strip()
            if not project or len(project) > 255:
                raise ValueError("loyiha nomi bo‘sh yoki juda uzun")
            sales = nonnegative(values[indexes["sales"]], "sales_plan")
            revenue = nonnegative(values[indexes["revenue"]], "revenue_plan")
            for metric, target in (("sales", sales), ("revenue", revenue)):
                row = db.query(V5BusinessPlan).filter_by(tenant_id=user.tenant_id, month=month, metric=metric, project=project).first()
                if row: row.target = target
                else: db.add(V5BusinessPlan(tenant_id=user.tenant_id, month=month, metric=metric, project=project, target=target))
            imported += 1
        if not imported:
            raise ValueError("Excel faylida ma’lumot qatori yo‘q")
        db.commit()
    except ValueError as exc:
        db.rollback()
        raise HTTPException(422, f"Excel qatori {number}: {exc}") from exc
    return {"imported": imported, "message": f"{imported} ta loyiha rejasi yuklandi"}


@router.get("")
def plan_fact(
    from_date: date = Query(alias="from"),
    to_date: date = Query(alias="to"),
    user=Depends(current_user),
    db: Session = Depends(get_db),
):
    if from_date > to_date:
        raise HTTPException(422, "Start date must not be after end date")
    if (to_date - from_date).days > 366:
        raise HTTPException(422, "Plan/fakt davri 366 kundan oshmasligi kerak")

    campaigns = tenant_query(db, Campaign, user).filter(
        Campaign.date >= datetime.combine(from_date, time.min, tzinfo=timezone.utc),
        Campaign.date <= datetime.combine(to_date, time.max, tzinfo=timezone.utc),
    ).all()
    leads = tenant_query(db, Lead, user).filter(
        Lead.created_at >= datetime.combine(from_date, time.min, tzinfo=timezone.utc),
        Lead.created_at <= datetime.combine(to_date, time.max, tzinfo=timezone.utc),
        Lead.crm == "bitrix",
    ).all()
    plans = tenant_query(db, V5BusinessPlan, user).filter(
        V5BusinessPlan.month >= from_date.strftime("%Y-%m"),
        V5BusinessPlan.month <= to_date.strftime("%Y-%m"),
    ).all()

    grouped_plans = defaultdict(lambda: defaultdict(lambda: Decimal("0")))
    for row in plans:
        metric = "sales" if row.metric in {"sales", "sale_count"} else "revenue" if row.metric in {"revenue", "sales_revenue"} else None
        if metric:
            grouped_plans[(row.month, metric)][row.project or "__total__"] += row.target
    def effective(month, metric):
        values = grouped_plans[(month, metric)]
        projects = [value for project, value in values.items() if project != "__total__"]
        return sum(projects, Decimal("0")) if projects else values.get("__total__", Decimal("0"))
    monthly_sales = defaultdict(lambda: Decimal("0"), {month: effective(month, "sales") for month, _ in grouped_plans if _ == "sales"})
    monthly_revenue = defaultdict(lambda: Decimal("0"), {month: effective(month, "revenue") for month, _ in grouped_plans if _ == "revenue"})
    period_days = [date.fromordinal(from_date.toordinal() + offset) for offset in range((to_date - from_date).days + 1)]
    plan_sales = sum((monthly_sales[day.strftime("%Y-%m")] / monthrange(day.year, day.month)[1] for day in period_days), Decimal("0"))
    plan_revenue = sum((monthly_revenue[day.strftime("%Y-%m")] / monthrange(day.year, day.month)[1] for day in period_days), Decimal("0"))
    won_leads = [row for row in leads if str(row.status).lower() in {"won", "paid", "sold", "success"}]
    fact_sales = sum(int(row.sales or 0) for row in campaigns) + len(won_leads)
    fact_revenue = sum((row.revenue or Decimal("0") for row in campaigns), Decimal("0")) + sum((row.amount or Decimal("0") for row in won_leads), Decimal("0"))
    total_days = len(period_days)
    today = datetime.now(timezone.utc).date()
    elapsed_days = min(total_days, max(0, (min(today, to_date) - from_date).days + 1)) if today >= from_date else 0
    forecast = round(fact_sales / elapsed_days * total_days) if elapsed_days else 0

    project_data = defaultdict(lambda: {"plan": Decimal("0"), "fact": 0, "plan_amount": Decimal("0"), "fact_amount": Decimal("0")})
    daily_data = defaultdict(int)
    for row in campaigns:
        project_data[row.name]["fact"] += int(row.sales or 0)
        project_data[row.name]["fact_amount"] += row.revenue or Decimal("0")
        if row.date:
            daily_data[row.date.date().isoformat()] += int(row.sales or 0)
    for row in won_leads:
        name = row.campaign_name or "Bitrix24"
        project_data[name]["fact"] += 1
        project_data[name]["fact_amount"] += row.amount or Decimal("0")
        if row.created_at:
            daily_data[row.created_at.date().isoformat()] += 1
    for row in plans:
        if not row.project or row.project == "__total__":
            continue
        year, month = map(int, row.month.split("-"))
        first, last = date(year, month, 1), date(year, month, monthrange(year, month)[1])
        overlap_start, overlap_end = max(first, from_date), min(last, to_date)
        if overlap_start > overlap_end:
            continue
        portion = Decimal((overlap_end - overlap_start).days + 1) / Decimal(monthrange(year, month)[1])
        if row.metric in {"sales", "sale_count"}: project_data[row.project]["plan"] += row.target * portion
        elif row.metric in {"revenue", "sales_revenue"}: project_data[row.project]["plan_amount"] += row.target * portion

    projects = []
    for name, values in sorted(project_data.items()):
        projects.append({
            "name": name,
            "plan": float(values["plan"]),
            "fact": values["fact"],
            "plan_amount": float(values["plan_amount"]),
            "fact_amount": float(values["fact_amount"]),
        })

    return {
        "start": from_date.isoformat(),
        "end": to_date.isoformat(),
        "plan": float(plan_sales),
        "fact": fact_sales,
        "plan_amount": float(plan_revenue),
        "fact_amount": float(fact_revenue),
        "forecast": forecast,
        "daily": [
            {"date": day.isoformat(), "plan": round(float(monthly_sales[day.strftime("%Y-%m")] / monthrange(day.year, day.month)[1]), 2), "fact": daily_data.get(day.isoformat(), 0)}
            for day in period_days
        ],
        "projects": projects,
    }
